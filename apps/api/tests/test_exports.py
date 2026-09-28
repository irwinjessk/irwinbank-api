from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from apps.api.tests.support import ApiTestCase
from apps.audit.models import JournalAudit

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def lire_xlsx(response):
    feuille = load_workbook(BytesIO(response.content)).active
    lignes = [list(ligne) for ligne in feuille.iter_rows(values_only=True)]
    return lignes


def tableau(lignes, premiere_colonne):
    """Renvoie (en-têtes, lignes de données) à partir de la ligne d'en-tête."""
    debut = next(i for i, ligne in enumerate(lignes) if ligne and ligne[0] == premiere_colonne)
    donnees = []
    for ligne in lignes[debut + 1:]:
        if not any(valeur is not None for valeur in ligne):
            break
        donnees.append(ligne)
    return lignes[debut], donnees


class ExportsTests(ApiTestCase):
    def setUp(self):
        super().setUp()
        self.banque = self.creer_banque()
        self.aya = self.creer_client(self.banque)
        self.compte = self.creer_compte(self.aya)
        for montant in ('100.00', '250.00'):
            self.client.post('/api/v1/transactions/', {'type_transaction': 'DEPOT', 'compte': self.compte.id, 'montant': montant}, format='json')
        self.client.post('/api/v1/transactions/', {'type_transaction': 'RETRAIT', 'compte': self.compte.id, 'montant': '30.00'}, format='json')

    def test_excel_des_transactions_filtre_par_type(self):
        response = self.client.get('/api/v1/transactions/export/?format=xlsx&type=DEPOT')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], XLSX)
        self.assertRegex(response['Content-Disposition'], r'attachment; filename="transactions-\d{4}-\d{2}-\d{2}-\d{4}\.xlsx"')
        lignes = lire_xlsx(response)
        self.assertEqual(lignes[0][0], 'Historique des transactions')
        self.assertTrue(any('Type = Dépôt' in (ligne[0] or '') for ligne in lignes[:5]))
        entetes, donnees = tableau(lignes, 'Date')
        self.assertEqual(entetes[:4], ['Date', 'Type', 'Sens', 'Montant'])
        self.assertEqual(sorted(Decimal(str(ligne[3])) for ligne in donnees), [Decimal('100'), Decimal('250')])
        self.assertTrue(all(ligne[1] == 'Dépôt' for ligne in donnees))
        resume = {ligne[0]: ligne[1] for ligne in lignes if ligne[0] in ('Total crédité', 'Total débité')}
        self.assertEqual(Decimal(str(resume['Total crédité'])), Decimal('350'))

    def test_pdf_des_transactions(self):
        response = self.client.get('/api/v1/transactions/export/?format=pdf&date_min=2020-01-01')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertIn('.pdf"', response['Content-Disposition'])

    def test_les_cinq_listes_s_exportent_dans_les_deux_formats(self):
        for ressource in ('banques', 'clients', 'comptes', 'transactions', 'factures'):
            for format_fichier, signature in (('pdf', b'%PDF'), ('xlsx', b'PK')):
                response = self.client.get(f'/api/v1/{ressource}/export/?format={format_fichier}')
                self.assertEqual(response.status_code, 200, f'{ressource} {format_fichier}')
                self.assertTrue(response.content.startswith(signature), f'{ressource} {format_fichier}')

    def test_export_vide_donne_un_fichier_avec_en_tetes(self):
        response = self.client.get('/api/v1/clients/export/?format=xlsx&nom=Personne')
        self.assertEqual(response.status_code, 200)
        entetes, donnees = tableau(lire_xlsx(response), 'Numéro')
        self.assertIn('E-mail', entetes)
        self.assertEqual(donnees, [])
        self.assertEqual(self.client.get('/api/v1/clients/export/?format=pdf&nom=Personne').status_code, 200)

    def test_agent_n_exporte_que_sa_banque(self):
        autre = self.creer_banque(nom='Autre')
        self.creer_client(autre, email='autre@example.com', nom='Diallo')
        self.connecter_agent(self.banque)
        lignes = lire_xlsx(self.client.get('/api/v1/clients/export/?format=xlsx'))
        _, donnees = tableau(lignes, 'Numéro')
        self.assertEqual([ligne[0] for ligne in donnees], [self.aya.numero_client])
        self.assertTrue(any((ligne[0] or '').startswith('Périmètre : ADA BANK') for ligne in lignes[:5]))

    def test_clients_archives_exclus_par_defaut(self):
        archive = self.creer_client(self.banque, email='old@example.com', nom='Ancien')
        archive.archive = True
        archive.save()
        _, actifs = tableau(lire_xlsx(self.client.get('/api/v1/clients/export/?format=xlsx')), 'Numéro')
        _, tous = tableau(lire_xlsx(self.client.get('/api/v1/clients/export/?format=xlsx&statut=tous')), 'Numéro')
        self.assertEqual((len(actifs), len(tous)), (1, 2))

    def test_export_trace_dans_le_journal(self):
        self.client.get('/api/v1/comptes/export/?format=pdf')
        ligne = JournalAudit.objects.get(action='export.genere')
        self.assertEqual(ligne.entite, 'comptes')
        self.assertIn('PDF', ligne.resume)

    def test_caracteres_hors_latin_ne_cassent_pas_le_pdf(self):
        self.creer_client(self.banque, email='nguyen@example.com', nom='Nguyễn 李 − “test”')
        response = self.client.get('/api/v1/clients/export/?format=pdf')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_format_inconnu_et_filtre_invalide(self):
        self.assertEqual(self.client.get('/api/v1/clients/export/?format=csv').status_code, 404)
        self.assertEqual(self.client.get('/api/v1/transactions/export/?format=xlsx&date_min=hier').status_code, 400)

    def test_sans_connexion(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get('/api/v1/clients/export/?format=pdf').status_code, 401)
