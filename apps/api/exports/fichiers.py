from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Callable
from xml.sax.saxutils import escape

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MARQUE = 'IRWIN BANK · ADA BANK (démo)'
COULEUR_ENTETE = '1E3A5F'
COULEUR_ZEBRE = 'F3F6FA'


@dataclass
class Colonne:
    titre: str
    valeur: Callable[[Any], Any]
    type: str = 'texte'  # texte | entier | montant | date | datetime
    largeur: float = 1.0


@dataclass
class Document:
    titre: str
    colonnes: list
    lignes: list
    meta: list
    resume: list


def _local(valeur):
    if isinstance(valeur, datetime):
        if timezone.is_aware(valeur):
            valeur = timezone.localtime(valeur)
        return valeur.replace(tzinfo=None)
    return valeur


def _montant_fr(valeur):
    return f'{Decimal(valeur):,.2f}'.replace(',', '\u00a0').replace('.', ',')


def texte_cellule(valeur, type_colonne):
    if valeur is None or valeur == '':
        return ''
    valeur = _local(valeur)
    if type_colonne == 'montant':
        return _montant_fr(valeur)
    if type_colonne == 'datetime' and isinstance(valeur, datetime):
        return valeur.strftime('%d/%m/%Y %H:%M')
    if type_colonne in ('date', 'datetime') and isinstance(valeur, (date, datetime)):
        return valeur.strftime('%d/%m/%Y')
    return str(valeur)


# ---------------------------------------------------------------- Excel


def generer_xlsx(document):
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = document.titre[:31]
    classeur.properties.creator = MARQUE
    classeur.properties.title = document.titre

    nb_colonnes = len(document.colonnes)
    feuille.cell(row=1, column=1, value=document.titre).font = Font(size=14, bold=True, color=COULEUR_ENTETE)
    for index, ligne_meta in enumerate(document.meta, start=2):
        feuille.cell(row=index, column=1, value=ligne_meta).font = Font(size=9, italic=True, color='666666')

    ligne_entete = len(document.meta) + 3
    bordure = Border(bottom=Side(style='thin', color='D0D7E2'))
    for numero, colonne in enumerate(document.colonnes, start=1):
        cellule = feuille.cell(row=ligne_entete, column=numero, value=colonne.titre)
        cellule.font = Font(bold=True, color='FFFFFF')
        cellule.fill = PatternFill('solid', fgColor=COULEUR_ENTETE)
        cellule.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    feuille.row_dimensions[ligne_entete].height = 22

    formats = {'montant': '#,##0.00', 'entier': '0', 'date': 'DD/MM/YYYY', 'datetime': 'DD/MM/YYYY HH:MM'}
    largeurs = [len(colonne.titre) + 2 for colonne in document.colonnes]
    for decalage, objet in enumerate(document.lignes, start=1):
        rangee = ligne_entete + decalage
        for numero, colonne in enumerate(document.colonnes, start=1):
            valeur = _local(colonne.valeur(objet))
            if colonne.type == 'montant' and valeur is not None:
                valeur = Decimal(valeur)
            cellule = feuille.cell(row=rangee, column=numero, value=valeur if valeur != '' else None)
            cellule.border = bordure
            if colonne.type in formats:
                cellule.number_format = formats[colonne.type]
            if colonne.type == 'montant':
                cellule.alignment = Alignment(horizontal='right')
            if decalage % 2 == 0:
                cellule.fill = PatternFill('solid', fgColor=COULEUR_ZEBRE)
            largeurs[numero - 1] = max(largeurs[numero - 1], len(texte_cellule(valeur, colonne.type)) + 2)

    derniere = ligne_entete + len(document.lignes)
    feuille.auto_filter.ref = f'A{ligne_entete}:{get_column_letter(nb_colonnes)}{max(derniere, ligne_entete)}'
    feuille.freeze_panes = feuille.cell(row=ligne_entete + 1, column=1)
    for numero, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(numero)].width = min(max(largeur, 8), 50)

    rangee = derniere + 2
    if not document.lignes:
        feuille.cell(row=rangee, column=1, value='Aucune ligne ne correspond aux filtres.').font = Font(italic=True, color='666666')
        rangee += 2
    for libelle, valeur in document.resume:
        feuille.cell(row=rangee, column=1, value=libelle).font = Font(bold=True)
        cellule = feuille.cell(row=rangee, column=2, value=Decimal(valeur) if isinstance(valeur, Decimal) else valeur)
        cellule.font = Font(bold=True)
        if isinstance(valeur, Decimal):
            cellule.number_format = formats['montant']
        rangee += 1

    feuille.page_setup.orientation = 'landscape'
    feuille.page_setup.fitToWidth = 1
    feuille.page_setup.fitToHeight = 0
    feuille.sheet_properties.pageSetUpPr.fitToPage = True
    feuille.print_title_rows = f'{ligne_entete}:{ligne_entete}'

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------- PDF


def _cp1252(texte):
    """Les polices PDF standard ne couvrent que le jeu Windows-1252."""
    texte = str(texte).replace('\u2212', '-').replace('\u202f', '\u00a0')
    return texte.encode('cp1252', 'replace').decode('cp1252')


def _pdf(texte):
    return escape(_cp1252(texte))


class _CanevasNumerote(pdf_canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._pages = []

    def showPage(self):
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._pages)
        for etat in self._pages:
            self.__dict__.update(etat)
            self._pied_de_page(total)
            super().showPage()
        super().save()

    def _pied_de_page(self, total):
        largeur, _ = self._pagesize
        self.setFont('Helvetica', 7.5)
        self.setFillColor(colors.HexColor('#6B7280'))
        self.drawString(1.5 * cm, 0.9 * cm, _cp1252(MARQUE))
        self.drawRightString(largeur - 1.5 * cm, 0.9 * cm, f'Page {self._pageNumber} / {total}')


def generer_pdf(document):
    tampon = BytesIO()
    marge = 1.5 * cm
    doc = SimpleDocTemplate(
        tampon,
        pagesize=landscape(A4),
        leftMargin=marge,
        rightMargin=marge,
        topMargin=marge,
        bottomMargin=1.6 * cm,
        title=document.titre,
        author=MARQUE,
    )
    style_titre = ParagraphStyle('titre', fontName='Helvetica-Bold', fontSize=16, leading=20, textColor=colors.HexColor(f'#{COULEUR_ENTETE}'))
    style_meta = ParagraphStyle('meta', fontName='Helvetica', fontSize=8.5, leading=11, textColor=colors.HexColor('#4B5563'))
    style_cellule = ParagraphStyle('cellule', fontName='Helvetica', fontSize=7.5, leading=9.5)
    style_montant = ParagraphStyle('montant', parent=style_cellule, alignment=TA_RIGHT)
    style_entete = ParagraphStyle('entete', parent=style_cellule, fontName='Helvetica-Bold', textColor=colors.white)
    style_resume = ParagraphStyle('resume', fontName='Helvetica-Bold', fontSize=9, leading=13)

    elements = [Paragraph(_pdf(document.titre), style_titre), Spacer(1, 4)]
    elements += [Paragraph(_pdf(ligne), style_meta) for ligne in document.meta]
    elements.append(Spacer(1, 10))

    style_entete_droite = ParagraphStyle('entete_droite', parent=style_entete, alignment=TA_RIGHT)
    donnees = [[
        Paragraph(_pdf(colonne.titre), style_entete_droite if colonne.type in ('montant', 'entier') else style_entete)
        for colonne in document.colonnes
    ]]
    for objet in document.lignes:
        donnees.append([
            Paragraph(
                _pdf(texte_cellule(colonne.valeur(objet), colonne.type)),
                style_montant if colonne.type in ('montant', 'entier') else style_cellule,
            )
            for colonne in document.colonnes
        ])

    total_poids = sum(colonne.largeur for colonne in document.colonnes)
    largeurs = [doc.width * colonne.largeur / total_poids for colonne in document.colonnes]
    tableau = Table(donnees, colWidths=largeurs, repeatRows=1)
    tableau.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(f'#{COULEUR_ENTETE}')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor(f'#{COULEUR_ZEBRE}')]),
        ('LINEBELOW', (0, 0), (-1, -1), 0.25, colors.HexColor('#D0D7E2')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(tableau)

    if not document.lignes:
        elements += [Spacer(1, 8), Paragraph('Aucune ligne ne correspond aux filtres.', style_meta)]
    if document.resume:
        elements.append(Spacer(1, 12))
        for libelle, valeur in document.resume:
            affichage = f'{_montant_fr(valeur)} F CFA' if isinstance(valeur, Decimal) else valeur
            elements.append(Paragraph(_pdf(f'{libelle} : {affichage}'), style_resume))

    doc.build(elements, canvasmaker=_CanevasNumerote)
    return tampon.getvalue()
